import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    BG_DARK = RGBColor(15, 23, 42)      # Slate 900
    CARD_BG = RGBColor(30, 41, 59)      # Slate 800
    TEXT_WHITE = RGBColor(248, 250, 252)# Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)# Slate 400
    ACCENT_CYAN = RGBColor(6, 182, 212) # Cyan 500
    ACCENT_GREEN = RGBColor(16, 185, 129)# Emerald 500

    def add_base_slide(title, category):
        slide = prs.slides.add_slide(blank_layout)
        
        # Background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()

        # Top Category / Module Pill
        cat_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(8), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_CYAN

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.65), Inches(12.1), Inches(0.6))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

        # Divider Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.3), Inches(12.133), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(51, 65, 85)
        line.line.fill.background()

        return slide

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = BG_DARK
    bg1.line.fill.background()

    # Hero Card
    hero = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.2), Inches(10.933), Inches(5.1))
    hero.fill.solid()
    hero.fill.fore_color.rgb = CARD_BG
    hero.line.color.rgb = RGBColor(51, 65, 85)

    tb = s1.shapes.add_textbox(Inches(1.8), Inches(1.6), Inches(9.733), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "DEVOPS CAPSTONE FINAL PRESENTATION"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_CYAN

    p1 = tf.add_paragraph()
    p1.text = "TaskBoard: Cloud-Native DevOps Architecture"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_before = Pt(10)

    p2 = tf.add_paragraph()
    p2.text = "End-to-End Delivery: FastAPI, React, CI/CD, DevSecOps, AWS EKS, Helm & Prometheus/Grafana"
    p2.font.size = Pt(16)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(8)

    p3 = tf.add_paragraph()
    p3.text = "\n• Author: Krushna Sonawane\n• GitHub Repository: github.com/AGENT-BABA/TaskBoard\n• Cloud Infrastructure: AWS EKS (Sydney ap-southeast-2)\n• Environment: Docker Compose, Kubernetes & GitOps"
    p3.font.size = Pt(14)
    p3.font.color.rgb = ACCENT_GREEN
    p3.space_before = Pt(14)

    # ==========================================
    # SLIDE 2: END-TO-END WORKFLOW ARCHITECTURE
    # ==========================================
    s2 = add_base_slide("End-to-End DevOps Lifecycle Architecture", "Overview • Full Delivery Pipeline")
    
    stages = [
        ("1. Development", "FastAPI + React + PostgreSQL\n• Alembic DB Migrations\n• Multi-stage Dockerfiles\n• Non-root runtime"),
        ("2. Testing (CI)", "Automated Pytest Suite\n• 8/8 tests passing\n• SQLite isolated test DB\n• Strict code quality gate"),
        ("3. DevSecOps", "GitHub Actions Pipeline\n• Trivy vulnerability scans\n• Starlette & Alpine patches\n• GHCR SHA image tagging"),
        ("4. Infrastructure", "Terraform AWS (IaC)\n• Custom VPC & Subnets\n• EKS Cluster & Node Group\n• Clean teardown ($0 leak)"),
        ("5. Kubernetes", "Helm Orchestration\n• 2x Pod High Availability\n• Ingress routing & HPA\n• ClusterIP services"),
        ("6. Observability", "Prometheus & Grafana\n• /metrics telemetry\n• Targets 100% UP\n• 4-panel Executive Dashboard\n• IaC auto-provisioning")
    ]

    card_w = Inches(3.8)
    card_h = Inches(2.6)
    for i, (stitle, sbody) in enumerate(stages):
        row = i // 3
        col = i % 3
        x = Inches(0.6 + col * 4.15)
        y = Inches(1.5 + row * 2.8)

        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = RGBColor(51, 65, 85)

        tb_c = s2.shapes.add_textbox(x + Inches(0.15), y + Inches(0.15), card_w - Inches(0.3), card_h - Inches(0.3))
        tfc = tb_c.text_frame
        tfc.word_wrap = True
        
        hp = tfc.paragraphs[0]
        hp.text = stitle
        hp.font.size = Pt(15)
        hp.font.bold = True
        hp.font.color.rgb = ACCENT_CYAN

        bp = tfc.add_paragraph()
        bp.text = sbody
        bp.font.size = Pt(12)
        bp.font.color.rgb = TEXT_WHITE
        bp.space_before = Pt(6)

    # Helper function for split content + LARGE image slide
    def add_split_slide(title, category, bullets, img_path, img_caption=""):
        slide = add_base_slide(title, category)
        
        # Left Text Card (Generous height, well-spaced)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.45), Inches(3.6), Inches(5.6))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = RGBColor(51, 65, 85)

        tb = slide.shapes.add_textbox(Inches(0.75), Inches(1.6), Inches(3.3), Inches(5.3))
        tf = tb.text_frame
        tf.word_wrap = True

        for idx, (bhead, btext) in enumerate(bullets):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = f"{bhead}\n"
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = ACCENT_CYAN
            if idx > 0:
                p.space_before = Pt(18)  # Generous vertical spacing fills the card

            run = p.add_run()
            run.text = btext
            run.font.bold = False
            run.font.size = Pt(11.5)
            run.font.color.rgb = TEXT_WHITE

        # Right Image (8.4 inches width)
        if os.path.exists(img_path):
            img_x = Inches(4.35)
            img_y = Inches(1.45)
            img_w = Inches(8.35)
            slide.shapes.add_picture(img_path, img_x, img_y, width=img_w)
            
            # Caption
            cap_b = slide.shapes.add_textbox(img_x, Inches(7.05), img_w, Inches(0.35))
            tfc = cap_b.text_frame
            pc = tfc.paragraphs[0]
            pc.alignment = PP_ALIGN.CENTER
            pc.text = img_caption or os.path.basename(img_path)
            pc.font.size = Pt(10)
            pc.font.color.rgb = TEXT_MUTED

        return slide

    # ==========================================
    # SLIDE 3: APPLICATION & DOCKER
    # ==========================================
    add_split_slide(
        "Application Architecture & Containerization",
        "Docker & Application Stack",
        [
            ("Modern Tech Stack", "FastAPI Python backend, React + Vite frontend, PostgreSQL 16 relational database with Alembic migrations."),
            ("Container Security Hardening", "Multi-stage Dockerfiles with separate build & runtime stages. Non-root user execution (`appuser:10001`)."),
            ("Race Condition Solved", "Solved PostgreSQL initdb startup delay using native Docker TCP healthcheck and `condition: service_healthy`."),
            ("Local Orchestration", "Single-command deployment with `docker compose up --build` connecting backend, frontend, postgres, and monitoring.")
        ],
        "docs/screenshots/01-TaskBoard-DashBoard.jpg",
        "Verified live TaskBoard UI running locally on localhost:3000"
    )

    # ==========================================
    # SLIDE 4: AUTOMATED TESTING
    # ==========================================
    add_split_slide(
        "Automated Testing & Code Quality Gates",
        "Pytest & SQLite Mock Isolation",
        [
            ("Pytest Suite", "8 comprehensive unit and integration tests covering CRUD operations, health endpoints, and payload validation."),
            ("Database Isolation", "Tests use an in-memory SQLite database (`conftest.py`), preventing contamination of production database."),
            ("Deterministic Results", "100% test pass rate executed in ~0.6 seconds before any container build or push."),
            ("CI Quality Gate", "Pytest is integrated into GitHub Actions — any failing test immediately blocks the CI/CD pipeline.")
        ],
        "docs/screenshots/02-pytest-passing.jpg",
        "Pytest terminal output showing 8/8 tests passing in SQLite test DB"
    )

    # ==========================================
    # SLIDE 5: CI/CD & DEVSECOPS
    # ==========================================
    add_split_slide(
        "CI/CD Pipeline & Container Security Scans",
        "GitHub Actions, Trivy & GHCR",
        [
            ("Automated Pipeline", "Triggered on push to `main` branch. Automates tests, builds, CVE scans, and image pushes."),
            ("DevSecOps with Trivy", "Vulnerability scans on backend & frontend images. Configured to fail on HIGH/CRITICAL severity."),
            ("CVE Remediation", "Patched Starlette CVE (upgraded to 0.45.3) and Alpine base vulnerabilities to achieve clean scans."),
            ("Immutable SHA Tagging", "Images published to GHCR tagged with Git commit SHA for full traceability and rollback capability.")
        ],
        "docs/screenshots/04-github-actions-success.jpg",
        "GitHub Actions workflow runs passing all build, test, and security jobs"
    )

    # ==========================================
    # SLIDE 6: TERRAFORM AWS INFRASTRUCTURE
    # ==========================================
    add_split_slide(
        "Infrastructure as Code: AWS EKS & VPC",
        "Terraform Cloud Infrastructure",
        [
            ("Terraform IaC", "Declarative HCL scripts provisioning full AWS infrastructure in Sydney (`ap-southeast-2`)."),
            ("VPC & Subnet Topology", "Custom VPC with multi-AZ public and private subnets, internet gateways, and routing tables."),
            ("Managed EKS Cluster", "Kubernetes cluster with managed EC2 `t3.small` worker nodes configured for cost-efficiency."),
            ("Enterprise SCP Compliance", "Configured `enable_irsa = false` and clean IAM roles to satisfy strict cloud boundary permissions.")
        ],
        "docs/screenshots/07-aws-eks-console.jpg",
        "AWS EKS Management Console confirming active taskboard-eks cluster"
    )

    # ==========================================
    # SLIDE 7: KUBERNETES & HELM DEPLOYMENT
    # ==========================================
    add_split_slide(
        "Kubernetes Orchestration & Helm Packaging",
        "Production Workloads & Routing",
        [
            ("Helm Packaging", "Clean templated Helm chart (`helm/taskboard`) managing Deployments, Services, Ingress, and HPA."),
            ("High Availability", "2x backend replicas and 2x frontend replicas deployed in isolated namespace `taskboard`."),
            ("Ingress Traffic Routing", "NGINX Ingress routing root path `/` to frontend and `/api` to backend microservice."),
            ("Dynamic Autoscaling", "HPA configured for CPU-based dynamic autoscaling under heavy user traffic loads.")
        ],
        "docs/screenshots/10-kubectl-pods-running.jpg",
        "Kubectl output confirming all backend, frontend, and database pods Running"
    )

    # ==========================================
    # SLIDE 8: OBSERVABILITY (PROMETHEUS)
    # ==========================================
    add_split_slide(
        "Observability: Metrics Telemetry & Prometheus",
        "Application Telemetry & Scraping",
        [
            ("Metrics Instrumentation", "FastAPI instrumented with `prometheus-fastapi-instrumentator` exposing `/metrics`."),
            ("Telemetry Captured", "Request rates, response duration histograms, process CPU, and memory utilization."),
            ("Prometheus Discovery", "Continuous automated scrape every 5s discovering `taskboard-backend`."),
            ("Target Health Verified", "Prometheus UI shows 100% active targets in `UP` state with zero scrape errors.")
        ],
        "docs/screenshots/15-prometheus-targets-01.jpg",
        "Prometheus Targets UI showing taskboard-backend endpoint healthy and UP"
    )

    # ==========================================
    # SLIDE 9: OBSERVABILITY (GRAFANA DASHBOARD)
    # ==========================================
    add_split_slide(
        "Executive Telemetry Dashboard (Grafana)",
        "Real-Time Visual Operations",
        [
            ("p95 Response Latency", "Radial gauge tracking 95th-percentile response time (`95 ms`) with green/orange/red SLA markers."),
            ("Total API Requests", "Live KPI stat card showing total processed requests (`324`) with an area trend sparkline."),
            ("Status Code Breakdown", "Donut chart detailing HTTP 2xx (100%) vs 4xx/5xx error rates with table metrics."),
            ("Throughput by Endpoint", "Multi-route time-series graph visualizing request rates across `/api/tasks`, `/health`, `/metrics`."),
            ("Infrastructure as Code", "Dashboard auto-provisioned automatically via Helm ConfigMap and Docker provisioning.")
        ],
        "docs/screenshots/16-grafana-dashboard-remake.jpg",
        "TaskBoard Executive Telemetry Dashboard featuring real-time multi-panel telemetry"
    )

    # ==========================================
    # SLIDE 10: CORE ENGINEERING CHALLENGES & ARCHITECTURE SOLUTIONS
    # ==========================================
    s10 = add_base_slide("Core Engineering Challenges & Architecture Solutions", "Technical Innovations & Problem-Solving")

    s10_cards = [
        ("1. PostgreSQL Initialization Race Condition Resolution (Zero Boot Failure)",
         "In multi-container stacks, Docker flags dependent containers as 'Started' upon process spawn before PostgreSQL finishes creating cluster files and opens TCP port 5432. This caused intermittent psycopg.OperationalError: connection refused crashes during initial Alembic migrations. Resolved by implementing a native Docker container TCP healthcheck (pg_isready -U taskboard) paired with condition: service_healthy, ensuring guaranteed deterministic startup across local and CI environments."),
        ("2. Proactive DevSecOps & Container CVE Remediation",
         "Standard container base images and dependencies often harbor latent vulnerabilities. Integrated Trivy vulnerability scanning directly into the GitHub Actions CI pipeline, enforcing automated build failures on HIGH or CRITICAL CVEs. Remediated Starlette CVEs by upgrading the dependency to 0.45.3 and hardened Alpine base runtime packages, producing clean, production-grade container artifacts published to GHCR with immutable Git commit SHA tags."),
        ("3. Universal GitOps Dashboard Provisioning (Docker + Kubernetes Parity)",
         "Custom telemetry dashboards created via Grafana UI are typically stored in transient container storage, risking missing metrics when evaluated on clean machines. Engineered a GitOps dashboard-as-code strategy: Docker Compose mounts automated provisioning configs, while the Helm chart deploys the dashboard JSON as a Kubernetes ConfigMap labeled grafana_dashboard: '1', auto-discovered by the Grafana sidecar on AWS EKS without manual intervention.")
    ]

    card_w = Inches(12.133)
    card_h = Inches(1.65)
    for idx, (htitle, hdesc) in enumerate(s10_cards):
        y_pos = Inches(1.5 + idx * 1.85)
        h_card = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), y_pos, card_w, card_h)
        h_card.fill.solid()
        h_card.fill.fore_color.rgb = CARD_BG
        h_card.line.color.rgb = RGBColor(51, 65, 85)

        tb_h = s10.shapes.add_textbox(Inches(0.85), y_pos + Inches(0.12), card_w - Inches(0.5), card_h - Inches(0.24))
        tf_h = tb_h.text_frame
        tf_h.word_wrap = True

        ph = tf_h.paragraphs[0]
        ph.text = f"{htitle}\n"
        ph.font.size = Pt(13)
        ph.font.bold = True
        ph.font.color.rgb = ACCENT_CYAN

        run_h = ph.add_run()
        run_h.text = hdesc
        run_h.font.bold = False
        run_h.font.size = Pt(11)
        run_h.font.color.rgb = TEXT_WHITE

    # ==========================================
    # SLIDE 11: CLOUD GOVERNANCE & PRODUCTION READINESS
    # ==========================================
    s11 = add_base_slide("Cloud Governance, Resilience & Capstone Takeaways", "Cloud Governance & Production Readiness")

    s11_cards = [
        ("1. Cloud Financial Governance & Verified Teardown ($0.00 Cost Leak)",
         "Managed cloud environments (EKS Control Plane, EC2 Worker Nodes, NAT Gateways, ELBs) incur continuous hourly infrastructure charges. Enforced strict Infrastructure-as-Code lifecycle management using Terraform. Following verified deployment testing and screenshot capture, executed an automated terraform destroy -auto-approve and verified zero remaining resources via the AWS CLI (ResourceNotFoundException), completely protecting cloud budgets."),
        ("2. Enterprise Kubernetes High Availability & Dynamic Autoscaling",
         "Packaged the microservice architecture into a modular Helm chart (helm/taskboard) featuring dual-replica backend and frontend deployments in dedicated namespace taskboard. Configured NGINX Ingress path-based routing (/ to React frontend, /api to FastAPI backend) and enabled Horizontal Pod Autoscaling (HPA) to automatically scale worker pods under dynamic CPU load thresholds."),
        ("3. End-to-End Enterprise DevOps Delivery Pipeline",
         "Successfully unified all isolated phases of modern software delivery—development, automated testing, containerization, DevSecOps scanning, cloud infrastructure provisioning, Kubernetes orchestration, and observability—into a reproducible, production-grade engineering workflow with complete auditability.")
    ]

    for idx, (htitle, hdesc) in enumerate(s11_cards):
        y_pos = Inches(1.5 + idx * 1.85)
        h_card = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), y_pos, card_w, card_h)
        h_card.fill.solid()
        h_card.fill.fore_color.rgb = CARD_BG
        h_card.line.color.rgb = RGBColor(51, 65, 85)

        tb_h = s11.shapes.add_textbox(Inches(0.85), y_pos + Inches(0.12), card_w - Inches(0.5), card_h - Inches(0.24))
        tf_h = tb_h.text_frame
        tf_h.word_wrap = True

        ph = tf_h.paragraphs[0]
        ph.text = f"{htitle}\n"
        ph.font.size = Pt(13)
        ph.font.bold = True
        ph.font.color.rgb = ACCENT_CYAN

        run_h = ph.add_run()
        run_h.text = hdesc
        run_h.font.bold = False
        run_h.font.size = Pt(11)
        run_h.font.color.rgb = TEXT_WHITE

    # Save Presentation to demo/
    os.makedirs("demo", exist_ok=True)
    output_path = os.path.join("demo", "TaskBoard_DevOps_Capstone_Presentation.pptx")
    try:
        prs.save(output_path)
        print(f"Presentation updated with 11 detailed slides: {output_path}")
    except PermissionError:
        alt_path = os.path.join("demo", "TaskBoard_DevOps_Capstone_Presentation_Updated.pptx")
        prs.save(alt_path)
        print(f"File locked by PowerPoint! Saved copy to: {alt_path}")

if __name__ == "__main__":
    create_deck()
